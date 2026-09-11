"""
services/document_reader_service.py

Servicio modular para la extracción y preparación de contenidos
de diversos formatos de archivo: Word, Excel, PowerPoint, Audio, Video,
Archivos comprimidos (ZIP/TAR) y Jupyter Notebooks.
"""

import os
import base64
import datetime
import json
import zipfile
import tarfile
import hashlib
import re
import glob
import tempfile
import shutil

def _find_soffice() -> str | None:
    """Busca la ruta del ejecutable soffice de LibreOffice en el sistema."""
    candidate = shutil.which("soffice")
    if candidate and os.path.exists(candidate):
        return candidate
    for path in [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]:
        if os.path.exists(path):
            return path
    try:
        import winreg
        k = winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\soffice.exe")
        val, _ = winreg.QueryValueEx(k, "")
        if val and os.path.exists(val):
            return val
    except Exception:
        pass
    return None

def _get_slide_num_from_name(fpath: str) -> int:
    m = re.search(r'\d+', os.path.basename(fpath))
    return int(m.group()) if m else 999999

# ── Límites de tamaño ──────────────────────────────────────────────────────────
MAX_MEDIA_BYTES    = 100 * 1024 * 1024  # 100 MB para audio/video
MAX_DOCX_BYTES     = 30 * 1024 * 1024   # 30 MB para documentos Word
MAX_EXCEL_BYTES    = 25 * 1024 * 1024   # 25 MB para hojas de cálculo
MAX_PPTX_BYTES     = 60 * 1024 * 1024   # 60 MB para presentaciones
MAX_ARCHIVE_BYTES  = 150 * 1024 * 1024  # 150 MB para inspeccionar ZIP/TAR
MAX_NOTEBOOK_BYTES = 20 * 1024 * 1024   # 20 MB para notebooks


def format_size(bytes_num: int) -> str:
    if bytes_num < 1024:
        return f"{bytes_num} B"
    elif bytes_num < 1024 * 1024:
        return f"{bytes_num / 1024:.1f} KB"
    else:
        return f"{bytes_num / (1024 * 1024):.1f} MB"


class DocumentReaderService:
    """Gestiona la lectura y transformación de documentos de diversa índole."""

    # ── Word (.docx) ─────────────────────────────────────────────────────────
    @staticmethod
    def read_docx(path: str, size: int) -> dict:
        if size > MAX_DOCX_BYTES:
            return {"type": "none", "message": f"Documento Word demasiado grande ({format_size(size)}). Límite: 30 MB."}
        try:
            import mammoth
            with open(path, "rb") as docx_file:
                res = mammoth.convert_to_html(docx_file)
                html_content = res.value
            return {
                "type": "docx",
                "html": html_content,
                "message": ""
            }
        except Exception as e:
            return {"type": "none", "message": f"Error al procesar documento Word: {str(e)}"}

    # ── Excel (.xlsx) ────────────────────────────────────────────────────────
    @staticmethod
    def read_excel(path: str, size: int) -> dict:
        if size > MAX_EXCEL_BYTES:
            return {"type": "none", "message": f"Archivo Excel demasiado grande ({format_size(size)}). Límite: 25 MB."}
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
            sheets_data = []

            for name in wb.sheetnames:
                ws = wb[name]
                rows_data = []
                # Limitamos a 400 filas y 40 columnas por hoja para rendimiento óptimo
                row_count = 0
                for row in ws.iter_rows(values_only=True):
                    row_count += 1
                    if row_count > 400:
                        break
                    # Limitar columnas y formatear valores
                    cleaned_row = []
                    for cell in list(row)[:40]:
                        if cell is None:
                            cleaned_row.append("")
                        elif isinstance(cell, datetime.datetime):
                            cleaned_row.append(cell.strftime("%Y-%m-%d %H:%M"))
                        elif isinstance(cell, datetime.date):
                            cleaned_row.append(cell.strftime("%Y-%m-%d"))
                        elif isinstance(cell, float):
                            cleaned_row.append(f"{cell:,.2f}" if cell.is_integer() is False else str(int(cell)))
                        else:
                            cleaned_row.append(str(cell))
                    rows_data.append(cleaned_row)

                sheets_data.append({
                    "name": name,
                    "rows": rows_data,
                    "has_more": row_count > 400
                })

            wb.close()
            return {
                "type": "excel",
                "sheets": sheets_data,
                "message": ""
            }
        except Exception as e:
            return {"type": "none", "message": f"Error al procesar hoja de cálculo Excel: {str(e)}"}

    # ── PowerPoint (.pptx) ───────────────────────────────────────────────────
    @staticmethod
    def read_pptx(path: str, size: int) -> dict:
        if size > MAX_PPTX_BYTES:
            return {"type": "none", "message": f"Presentación PowerPoint demasiado grande ({format_size(size)}). Límite: 60 MB."}

        abs_path = os.path.abspath(path)
        try:
            mtime = os.path.getmtime(abs_path)
        except OSError:
            mtime = 0

        # Clave y directorio de caché únicos basados en ruta, modificación y tamaño
        cache_key = hashlib.md5(f"{abs_path}_{mtime}_{size}".encode("utf-8")).hexdigest()
        cache_dir = os.path.join(tempfile.gettempdir(), "explorador_fb_cache", "pptx", cache_key)
        os.makedirs(cache_dir, exist_ok=True)

        # ── 1. Comprobar si ya existen diapositivas renderizadas en caché ────
        cached_pngs = [
            f for f in glob.glob(os.path.join(cache_dir, "*.*"))
            if f.lower().endswith(('.png', '.jpg', '.jpeg'))
        ]
        if cached_pngs:
            cached_pngs.sort(key=_get_slide_num_from_name)
            return DocumentReaderService._build_visual_pptx_response(cached_pngs)

        # ── 2. Método A: PowerPoint COM (Fidelidad 100 % idéntica al original) ───
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            try:
                ppt_app = win32com.client.Dispatch("PowerPoint.Application")
                # WithWindow = 0 (msoFalse) ejecuta la conversión en segundo plano sin desplegar ventanas
                pres = ppt_app.Presentations.Open(abs_path, -1, 0, 0)
                # 18 = ppSaveAsPNG
                pres.SaveAs(cache_dir, 18)
                pres.Close()
                ppt_app.Quit()
            finally:
                pythoncom.CoUninitialize()

            pngs = [
                f for f in glob.glob(os.path.join(cache_dir, "*.*"))
                if f.lower().endswith(('.png', '.jpg', '.jpeg'))
            ]
            if pngs:
                pngs.sort(key=_get_slide_num_from_name)
                return DocumentReaderService._build_visual_pptx_response(pngs)
        except Exception:
            pass

        # ── 3. Método B: LibreOffice headless + pypdfium2 ────────────────────
        soffice_bin = _find_soffice()
        if soffice_bin:
            try:
                import subprocess
                temp_pdf_dir = os.path.join(tempfile.gettempdir(), "explorador_fb_cache", "tmp_pdf_" + cache_key)
                os.makedirs(temp_pdf_dir, exist_ok=True)
                cmd = [soffice_bin, "--headless", "--convert-to", "pdf", "--outdir", temp_pdf_dir, abs_path]
                subprocess.run(cmd, capture_output=True, text=True, timeout=40)
                base_name = os.path.splitext(os.path.basename(abs_path))[0]
                pdf_file = os.path.join(temp_pdf_dir, base_name + ".pdf")
                if os.path.exists(pdf_file):
                    import pypdfium2
                    pdf_doc = pypdfium2.PdfDocument(pdf_file)
                    for p_idx in range(len(pdf_doc)):
                        page = pdf_doc[p_idx]
                        img = page.render(scale=1.5).to_pil()
                        slide_path = os.path.join(cache_dir, f"Diapositiva{p_idx + 1}.png")
                        img.save(slide_path, format="PNG")
                    try:
                        os.remove(pdf_file)
                        os.rmdir(temp_pdf_dir)
                    except Exception:
                        pass
                    pngs = [
                        f for f in glob.glob(os.path.join(cache_dir, "*.*"))
                        if f.lower().endswith(('.png', '.jpg', '.jpeg'))
                    ]
                    if pngs:
                        pngs.sort(key=_get_slide_num_from_name)
                        return DocumentReaderService._build_visual_pptx_response(pngs)
            except Exception:
                pass

        # ── 4. Método C: python-pptx (Respaldo estructurado de textos y formas) ──
        try:
            import pptx
            prs = pptx.Presentation(abs_path)
            slides_data = []

            for idx, slide in enumerate(prs.slides):
                slide_info = {
                    "number": idx + 1,
                    "title": "",
                    "paragraphs": [],
                    "tables": [],
                    "images": []
                }

                try:
                    if slide.shapes.title and slide.shapes.title.has_text_frame:
                        slide_info["title"] = slide.shapes.title.text_frame.text.strip()
                except Exception:
                    pass

                for shape in slide.shapes:
                    if shape == getattr(slide.shapes, "title", None):
                        continue
                    if shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            txt = p.text.strip()
                            if txt:
                                slide_info["paragraphs"].append({
                                    "text": txt,
                                    "level": p.level
                                })
                    if shape.has_table:
                        tbl = shape.table
                        tbl_rows = []
                        for row in tbl.rows:
                            tbl_rows.append([c.text.strip() for c in row.cells])
                        slide_info["tables"].append(tbl_rows)
                    if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                        try:
                            blob = shape.image.blob
                            ct = shape.image.content_type or "image/png"
                            b64 = base64.b64encode(blob).decode("utf-8")
                            slide_info["images"].append(f"data:{ct};base64,{b64}")
                        except Exception:
                            pass

                slides_data.append(slide_info)

            return {
                "type": "pptx",
                "mode": "fallback",
                "slides": slides_data,
                "total_slides": len(slides_data),
                "message": ""
            }
        except Exception as e:
            return {"type": "none", "message": f"Error al procesar presentación PowerPoint: {str(e)}"}

    @staticmethod
    def _build_visual_pptx_response(png_paths: list[str]) -> dict:
        slides = []
        for idx, p in enumerate(png_paths):
            try:
                with open(p, "rb") as f:
                    b64 = base64.b64encode(f.read()).decode("utf-8")
                slides.append({
                    "number": idx + 1,
                    "image": f"data:image/png;base64,{b64}",
                    "title": f"Diapositiva {idx + 1}"
                })
            except Exception:
                continue
        return {
            "type": "pptx",
            "mode": "visual",
            "slides": slides,
            "total_slides": len(slides),
            "message": ""
        }

    # ── Archivos comprimidos (.zip, .tar, .gz) ───────────────────────────────
    @staticmethod
    def read_archive(path: str, ext: str, size: int) -> dict:
        if size > MAX_ARCHIVE_BYTES:
            return {"type": "none", "message": f"Archivo comprimido demasiado grande ({format_size(size)}). Límite: 150 MB."}

        entries = []
        total_uncompressed = 0
        total_compressed = 0

        try:
            if ext in {'.zip'}:
                with zipfile.ZipFile(path, 'r') as z:
                    for info in z.infolist():
                        total_uncompressed += info.file_size
                        total_compressed += info.compress_size
                        ratio = 0
                        if info.file_size > 0:
                            ratio = round((1 - (info.compress_size / info.file_size)) * 100)
                        date_str = "%04d-%02d-%02d %02d:%02d" % info.date_time[:5]
                        entries.append({
                            "name": info.filename,
                            "size": info.file_size,
                            "size_fmt": format_size(info.file_size),
                            "compressed": info.compress_size,
                            "compressed_fmt": format_size(info.compress_size),
                            "ratio": f"{max(0, ratio)} %",
                            "is_dir": info.is_dir(),
                            "date": date_str
                        })
            elif ext in {'.tar', '.gz', '.tgz', '.bz2'}:
                mode = 'r:*'
                with tarfile.open(path, mode) as t:
                    for member in t.getmembers():
                        total_uncompressed += member.size
                        date_str = datetime.datetime.fromtimestamp(member.mtime).strftime("%Y-%m-%d %H:%M")
                        entries.append({
                            "name": member.name,
                            "size": member.size,
                            "size_fmt": format_size(member.size),
                            "compressed": member.size,
                            "compressed_fmt": format_size(member.size),
                            "ratio": "0 %",
                            "is_dir": member.isdir(),
                            "date": date_str
                        })

            return {
                "type": "archive",
                "entries": entries,
                "total_files": len(entries),
                "total_uncompressed_fmt": format_size(total_uncompressed),
                "total_compressed_fmt": format_size(total_compressed if total_compressed > 0 else size),
                "message": ""
            }
        except Exception as e:
            return {"type": "none", "message": f"Error al leer archivo comprimido: {str(e)}"}

    # ── Audio y Video ────────────────────────────────────────────────────────
    @staticmethod
    def read_media(path: str, ext: str, size: int, media_kind: str) -> dict:
        if size > MAX_MEDIA_BYTES:
            return {"type": "none", "message": f"Archivo multimedia demasiado grande ({format_size(size)}). Límite: 100 MB."}

        mime_types = {
            '.mp3': 'audio/mpeg', '.wav': 'audio/wav', '.ogg': 'audio/ogg',
            '.m4a': 'audio/mp4', '.aac': 'audio/aac', '.flac': 'audio/flac',
            '.mp4': 'video/mp4', '.webm': 'video/webm', '.mov': 'video/quicktime',
        }
        mime = mime_types.get(ext, f"{media_kind}/{ext.lstrip('.')}")

        try:
            with open(path, "rb") as f:
                data = base64.b64encode(f.read()).decode("utf-8")
            return {
                "type": media_kind,
                "data_url": f"data:{mime};base64,{data}",
                "mime": mime,
                "message": ""
            }
        except Exception as e:
            return {"type": "none", "message": f"Error al cargar archivo multimedia: {str(e)}"}

    # ── Jupyter Notebooks (.ipynb) ───────────────────────────────────────────
    @staticmethod
    def read_notebook(path: str, size: int) -> dict:
        if size > MAX_NOTEBOOK_BYTES:
            return {"type": "none", "message": f"Cuaderno Jupyter demasiado grande ({format_size(size)}). Límite: 20 MB."}
        try:
            with open(path, "r", encoding="utf-8") as f:
                nb = json.load(f)
            cells = []
            for cell in nb.get("cells", []):
                ctype = cell.get("cell_type", "")
                src = "".join(cell.get("source", []))
                outputs = []
                for out in cell.get("outputs", []):
                    # Salida de texto
                    if "text" in out:
                        outputs.append({"type": "text", "content": "".join(out["text"])})
                    elif "data" in out:
                        data_dict = out["data"]
                        if "image/png" in data_dict:
                            outputs.append({"type": "image", "content": f"data:image/png;base64,{data_dict['image/png']}"})
                        elif "text/plain" in data_dict:
                            outputs.append({"type": "text", "content": "".join(data_dict["text/plain"])})
                    elif out.get("output_type") == "error":
                        outputs.append({"type": "error", "content": "\n".join(out.get("traceback", []))})

                cells.append({
                    "type": ctype,
                    "source": src,
                    "outputs": outputs
                })

            return {
                "type": "notebook",
                "cells": cells,
                "message": ""
            }
        except Exception as e:
            return {"type": "none", "message": f"Error al procesar cuaderno Jupyter: {str(e)}"}
